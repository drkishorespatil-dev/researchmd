#!/usr/bin/env python3
"""Build the blinded, Cureus-compliant manuscript."""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ---------- base styling ----------
st = doc.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
pf = st.paragraph_format
pf.space_after = Pt(8)
pf.line_spacing = 1.3

for name, size in (('Heading 1', 14), ('Heading 2', 12.5)):
    h = doc.styles[name]
    h.font.name = 'Times New Roman'
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.color.rgb = RGBColor(0, 0, 0)
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)

for s in doc.sections:
    s.top_margin = s.bottom_margin = Inches(1)
    s.left_margin = s.right_margin = Inches(1)

ITAL = re.compile(r'(\*[^*]+\*)')


def runs(par, text, bold=False, italic=False, color=None, size=None):
    """Add text to a paragraph, honouring *italic* markers."""
    for chunk in ITAL.split(text):
        if not chunk:
            continue
        it = italic
        if chunk.startswith('*') and chunk.endswith('*') and len(chunk) > 2:
            chunk = chunk[1:-1]
            it = True
        r = par.add_run(chunk)
        r.bold = bold
        r.italic = it
        if color:
            r.font.color.rgb = color
        if size:
            r.font.size = Pt(size)
    return par


def p(text='', bold=False, italic=False, align=None, size=None, color=None,
      space_after=None, space_before=None):
    par = doc.add_paragraph()
    runs(par, text, bold=bold, italic=italic, color=color, size=size)
    if align is not None:
        par.alignment = align
    if space_after is not None:
        par.paragraph_format.space_after = Pt(space_after)
    if space_before is not None:
        par.paragraph_format.space_before = Pt(space_before)
    return par


def h1(text):
    par = doc.add_heading(level=1)
    runs(par, text, bold=True)
    return par


def h2(text):
    par = doc.add_paragraph()
    runs(par, text, bold=True, italic=True)
    par.paragraph_format.space_before = Pt(12)
    par.paragraph_format.space_after = Pt(4)
    return par


QUERY = RGBColor(0xC0, 0x00, 0x00)


def query(text):
    par = doc.add_paragraph()
    runs(par, '[AUTHOR QUERY - resolve and delete before submission] ' + text,
         bold=True, color=QUERY, size=11)
    return par


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear')
    el.set(qn('w:color'), 'auto')
    el.set(qn('w:fill'), hexcolor)
    tcPr.append(el)


def table(title, rows, legend, bold_last_row=True):
    """title above, table, legend below - Cureus convention."""
    tp = doc.add_paragraph()
    runs(tp, title, bold=True)
    tp.paragraph_format.space_before = Pt(12)
    tp.paragraph_format.space_after = Pt(4)

    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ''
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.space_before = Pt(2)
            par.paragraph_format.line_spacing = 1.0
            is_hdr = (i == 0)
            is_tot = bold_last_row and i == len(rows) - 1
            runs(par, val, bold=is_hdr or is_tot, size=11)
            if j > 0:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if is_hdr:
                shade(cell, 'D9D9D9')
            elif is_tot:
                shade(cell, 'F2F2F2')

    lp = doc.add_paragraph()
    runs(lp, legend, size=10.5)
    lp.paragraph_format.space_before = Pt(4)
    lp.paragraph_format.space_after = Pt(12)


def figure(title, legend, placeholder):
    box = doc.add_table(rows=1, cols=1)
    box.style = 'Table Grid'
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = box.cell(0, 0)
    cell.text = ''
    par = cell.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(28)
    par.paragraph_format.space_after = Pt(28)
    runs(par, placeholder, bold=True, color=QUERY, size=11)
    shade(cell, 'FAFAFA')

    cp = doc.add_paragraph()
    runs(cp, title, bold=True, size=11)
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(6)
    cp.paragraph_format.space_after = Pt(2)

    lp = doc.add_paragraph()
    runs(lp, legend, size=10.5)
    lp.paragraph_format.space_after = Pt(12)


def rule():
    par = doc.add_paragraph()
    pPr = par._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '808080')
    pbdr.append(bottom)
    pPr.append(pbdr)
    par.paragraph_format.space_after = Pt(10)


# ==================== FRONT MATTER ====================
p('BLINDED MANUSCRIPT - ALL AUTHOR AND INSTITUTIONAL IDENTIFIERS REMOVED',
  bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
rule()

p('Article type: Original Article', size=11)
p('Study design (submission questionnaire): Prospective observational study', size=11)
p('Categories: General Surgery; Gastroenterology', size=11)
p('Authors and affiliations: [REDACTED FOR BLINDED REVIEW - five authors, two departments, '
  'one institution. Restore the full author table, with affiliation strings character-identical '
  'across authors sharing an institution, in the unblinded submission copy.]', size=11)
p('Keywords: cholelithiasis; upper gastrointestinal endoscopy; laparoscopic cholecystectomy; '
  'gastritis; post-cholecystectomy syndrome; dyspepsia; preoperative evaluation; '
  'esophagogastroduodenoscopy', size=11)
rule()

p('Title', bold=True)
p('Routine Upper Gastrointestinal Endoscopy in Patients With Cholelithiasis Presenting With '
  'Upper Gastrointestinal Symptoms: A Prospective Observational Study From a Tertiary Care '
  'Center in Central India', bold=True)

# ==================== ABSTRACT ====================
h1('ABSTRACT')

h2('Background')
p('Upper gastrointestinal symptoms in patients with gallstones overlap substantially with those '
  'of mucosal disease of the esophagus, stomach, and duodenum. Because ultrasonography cannot '
  'resolve this overlap, some patients undergo cholecystectomy for symptoms that do not arise '
  'from the gallbladder and remain symptomatic afterwards. Whether upper gastrointestinal '
  'endoscopy should routinely precede cholecystectomy remains unsettled.')

h2('Methodology')
p('Eighty-five consecutive adults with ultrasonographically confirmed cholelithiasis and at '
  'least one upper gastrointestinal symptom were enrolled prospectively at a government tertiary '
  'care center in central India. All underwent esophagogastroduodenoscopy before operation, '
  'cholecystectomy, and repeat esophagogastroduodenoscopy with symptom reassessment one month '
  'after operation. Lesions detected before operation were treated medically alongside surgery. '
  "Associations were tested using Fisher's exact test and the Fisher-Freeman-Halton exact test.")

h2('Results')
p('Mean age was 45.96 years (standard deviation 17.74) and 60 patients (70.6%) were female. '
  'Epigastric pain was universal; nausea (51.8%), regurgitation (48.2%), postprandial fullness '
  '(47.1%), bloating (43.5%), and vomiting (40.0%) were variably present. Preoperative endoscopy '
  'was abnormal in 75 patients (88.2%): gastritis in 46 (54.1%), duodenitis in 12 (14.1%), '
  'esophagitis in 11 (12.9%), and peptic ulcer in six (7.1%). Neither age, sex, nor symptom '
  'duration predicted an abnormal preoperative endoscopy (p = 0.965, 0.716, and 0.855 '
  'respectively). One month after operation, 64 patients (75.3%) were symptom-free and 21 '
  '(24.7%) reported persistent symptoms. Symptom resolution was significantly more frequent '
  'where preoperative endoscopy had been abnormal and treated (60/75, 80.0%) than where it had '
  'been normal (4/10, 40.0%), with an odds ratio of 6.00 (95% confidence interval 1.50-23.99, '
  'p = 0.013). Among the 75 patients with abnormal preoperative findings, repeat endoscopy '
  'improved in 60 (80.0%), was unchanged in 11 (14.7%), and had worsened in four (5.3%), and '
  'endoscopic outcome was significantly associated with preoperative endoscopic status '
  '(p < 0.001). Persistent symptoms tracked residual mucosal disease closely: 12 of the 15 '
  'patients (80.0%) whose repeat endoscopy remained abnormal were still symptomatic, against '
  'nine of 70 (12.9%) whose repeat endoscopy was normal (odds ratio 27.11, 95% confidence '
  'interval 6.39-115.10, p < 0.001).')

h2('Conclusions')
p('Coexisting upper gastrointestinal mucosal disease was present in almost 90% of patients with '
  'symptomatic gallstones and proved largely reversible once identified and treated. A normal '
  'preoperative endoscopy identified the patients least likely to benefit symptomatically from '
  'cholecystectomy, and symptoms persisting afterwards tracked residual mucosal disease rather '
  'than the biliary operation. Preoperative endoscopy discloses a treatable and otherwise occult '
  'source of symptoms and merits a low threshold for use where the background burden of '
  'gastroduodenal disease is high.')

# ==================== INTRODUCTION ====================
h1('INTRODUCTION')
p('Gallstone disease is among the most common disorders of the digestive tract worldwide. A '
  'systematic review and meta-analysis of studies published this century estimated a pooled '
  'global prevalence of 6.1%, consistently higher in women and rising with age [1]. In the '
  'United States, population-level data show both a substantial and an increasing burden, with '
  'prevalence rising from 7.4% in 1988-1994 to 13.9% in 2017-2020 [2]. Stone formation reflects '
  'the interaction of biliary cholesterol supersaturation, accelerated crystal nucleation, mucin '
  'accumulation, and gallbladder hypomotility, modified by genetic susceptibility and by '
  'metabolic factors such as insulin resistance, obesity, and dyslipidemia [3,4]. Estrogen '
  'exposure is a further determinant, and menopausal hormone therapy has been shown in a large '
  'population-based cohort to increase gallstone risk across treatment types [5].')
p('Laparoscopic cholecystectomy is the definitive treatment for symptomatic gallstones, and '
  'current guidelines position abdominal ultrasonography as the first-line investigation [6,7]. '
  'Ultrasonography, however, answers only one question. It establishes whether stones are '
  'present; it says nothing about the mucosa of the esophagus, stomach, or duodenum. This matters '
  'because the symptoms that bring patients to operation are frequently not the classical, '
  'well-localized, self-limiting attacks of biliary colic, but a diffuse complex of epigastric '
  'discomfort, nausea, bloating, postprandial fullness, and regurgitation. That complex is '
  'equally characteristic of gastritis, peptic ulcer disease, gastroesophageal reflux, and '
  'functional dyspepsia. When a patient with such symptoms is found to have gallstones, the '
  'stones are assumed to be the cause, and the assumption is often untested.')
p('The consequence is visible in outcomes after operation. A substantial minority of patients '
  'continue to experience upper gastrointestinal symptoms after cholecystectomy, and the '
  'symptomatic benefit of the operation is least predictable precisely in those whose '
  'preoperative symptoms were atypical [8]. Post-cholecystectomy syndrome, the label applied to '
  'this group, is not a single disease but a heterogeneous collection of biliary and non-biliary '
  'disorders, many of them treatable once identified [9]. Several of the commonest contributors, '
  'including reflux disease, peptic ulcer disease, and gastritis, are diagnosable before '
  'operation by endoscopy and treatable by medical therapy alone.')
p('Against this background, several groups have asked whether esophagogastroduodenoscopy should '
  'precede cholecystectomy. Reported yields vary widely. Rashid et al. found abnormalities in 35% '
  'of routinely scoped patients and reported that symptom persistence after operation fell from '
  '32.8% to 3.3% when endoscopy was performed and lesions treated [10]. Morrison and Mokoena '
  'documented abnormal findings in 35.4% of patients and active pathology requiring treatment '
  'before operation in 28.2% [11]. Considerably higher yields have been reported from South Asian '
  'centers, with abnormal findings in 75.5% of 400 patients in one series, whose authors '
  'nonetheless concluded that routine endoscopy is not required for every patient and is of '
  'greatest value in those with atypical presentations [12]. The disagreement is not merely '
  'methodological: the prior probability of finding treatable mucosal disease depends on the '
  'background prevalence of gastroduodenal pathology in the population being scoped, which in '
  'parts of India is high, reflecting the local burden of *Helicobacter pylori* infection [13].')
p('Data from Indian government tertiary care centers, which serve predominantly low- and '
  'middle-income populations for whom a second episode of investigation carries real economic '
  'cost, remain limited. This study was undertaken to determine the prevalence and spectrum of '
  'upper gastrointestinal endoscopic abnormalities in patients with cholelithiasis presenting '
  'with upper gastrointestinal symptoms, and to establish whether preoperative endoscopic '
  'findings, and their response to treatment, relate to symptomatic and endoscopic outcome after '
  'cholecystectomy.')

# ==================== METHODS ====================
h1('MATERIALS AND METHODS')

h2('Study design and setting')
p('This prospective observational study was conducted in the Department of General Surgery of a '
  'government autonomous tertiary care teaching hospital in central India. The protocol was '
  'reviewed and approved by the institutional Scientific Review Committee and Human Research '
  'Ethics Committee before enrollment began. Written informed consent was obtained from every '
  'participant after the participant information sheet had been explained in a language the '
  'participant understood.')

h2('Participants')
p('Consecutive patients aged 18 years or older with cholelithiasis confirmed on abdominal '
  'ultrasonography, presenting with one or more of upper abdominal discomfort, nausea or '
  'vomiting, or abdominal bloating or fullness, were eligible. Patients were excluded if they had '
  'complicated gallstone disease, defined as acute or chronic cholecystitis, choledocholithiasis, '
  'cholangitis, gallstone pancreatitis, empyema of the gallbladder, gallbladder polyp, '
  'gallbladder adenomyomatosis, or carcinoma of the gallbladder; if they had undergone previous '
  'biliary, pancreatic, or other abdominal surgery; or if they declined participation. '
  'Convenience sampling was used.')
p('The sample size was derived from the finding of Kunnuru et al., in which preoperative '
  'endoscopy was abnormal in 75.5% of patients with symptomatic cholelithiasis [12]. Using '
  'n = z²pq/d² with p = 75.5%, 95% confidence, and 10% absolute precision, the minimum '
  'required sample was 72. Eighty-five consecutive eligible patients were enrolled during the '
  'study period, exceeding this minimum.')

h2('Procedures')
p('All participants underwent a structured history and clinical examination recorded on a '
  'pre-designed, pre-tested semi-structured proforma, together with routine hematological and '
  'biochemical investigations and abdominal ultrasonography. Six upper gastrointestinal symptoms '
  'were recorded as present or absent: epigastric pain, nausea, vomiting, bloating, postprandial '
  'fullness, and regurgitation. The duration of symptoms before presentation was recorded in '
  'weeks.')
p('Every participant then underwent diagnostic esophagogastroduodenoscopy before operation using '
  'a flexible video gastroscope (Olympus Corporation, Tokyo, Japan). Findings were assigned to a '
  'single predominant category: gastritis, duodenitis, esophagitis, peptic ulcer, or normal. '
  'Patients in whom mucosal pathology was identified received appropriate medical therapy in '
  'addition to surgical treatment. After preoperative optimization and anesthetic clearance, '
  'cholecystectomy was performed by the laparoscopic approach, with conversion to open '
  'cholecystectomy where dictated by intraoperative findings.')
p('Participants were reassessed one month after operation. The same six symptoms were re-elicited '
  'and esophagogastroduodenoscopy was repeated in every participant. The postoperative endoscopic '
  'outcome was classified relative to the preoperative examination as improved, unchanged, or '
  'worsened; for participants whose preoperative examination had been normal, a normal repeat '
  'examination was classified as unchanged, since improvement from a normal baseline is not '
  'definable. A participant was recorded as symptom-free if none of the six index symptoms was '
  'present at the one-month assessment.')

h2('Statistical analysis')
p('Data were entered in Microsoft Excel (Microsoft Corporation, Redmond, WA, USA) and analyzed '
  'using IBM SPSS Statistics for Windows, Version 26 (IBM Corp., Armonk, NY, USA). Continuous '
  'variables are summarized as mean and standard deviation, categorical variables as frequency '
  'and percentage. Independent-samples t-tests compared continuous variables between endoscopy '
  "groups. Associations in two-by-two tables were tested using Fisher's exact test, with odds "
  'ratios and 95% confidence intervals calculated by the Woolf method; the '
  'Fisher-Freeman-Halton exact test was used for larger contingency tables. A two-tailed p value '
  'below 0.05 was considered statistically significant.')

# ==================== RESULTS ====================
h1('RESULTS')

h2('Demographic profile')
p('Eighty-five patients were enrolled and all completed the protocol, including postoperative '
  'endoscopy. Age ranged from 19 to 77 years, with a mean of 45.96 years (standard deviation '
  '17.74) and a median of 46 years. Patients aged 41-60 years formed the largest group. Sixty '
  'patients (70.6%) were female and 25 (29.4%) male, a female-to-male ratio of 2.4:1. The age '
  'distribution is shown in Table 1.')

table('TABLE 1: Age distribution of the study population',
      [['Age group (years)', 'Frequency', 'Percentage (%)'],
       ['≤20', '5', '5.9'],
       ['21-40', '28', '32.9'],
       ['41-60', '31', '36.5'],
       ['≥61', '21', '24.7'],
       ['Total', '85', '100.0']],
      'Mean age 45.96 years (standard deviation 17.74), median 46 years, range 19-77 years.')

h2('Clinical presentation')
p('Epigastric pain was present in every patient. Nausea was the next most frequent symptom, '
  'followed by regurgitation, postprandial fullness, bloating, and vomiting. The mean duration of '
  'symptoms before presentation was 9.21 weeks. The distribution of presenting symptoms is shown '
  'in Table 2.')

table('TABLE 2: Distribution of presenting upper gastrointestinal symptoms',
      [['Symptom', 'Present, n (%)', 'Absent, n (%)'],
       ['Epigastric pain', '85 (100.0)', '0 (0.0)'],
       ['Nausea', '44 (51.8)', '41 (48.2)'],
       ['Regurgitation', '41 (48.2)', '44 (51.8)'],
       ['Postprandial fullness', '40 (47.1)', '45 (52.9)'],
       ['Bloating', '37 (43.5)', '48 (56.5)'],
       ['Vomiting', '34 (40.0)', '51 (60.0)']],
      'Symptoms are not mutually exclusive; each patient could report more than one. n = 85.',
      bold_last_row=False)

h2('Preoperative endoscopic findings')
p('Preoperative esophagogastroduodenoscopy was abnormal in 75 patients (88.2%) and normal in 10 '
  '(11.8%). Gastritis was the dominant lesion, identified in 46 patients (54.1%), followed by '
  'duodenitis in 12 (14.1%), esophagitis in 11 (12.9%), and peptic ulcer in six (7.1%). Findings '
  'are shown in Table 3, and representative endoscopic appearances in Figures 1-3.')

table('TABLE 3: Preoperative upper gastrointestinal endoscopic findings',
      [['Finding', 'Frequency', 'Percentage (%)'],
       ['Gastritis', '46', '54.1'],
       ['Duodenitis', '12', '14.1'],
       ['Esophagitis', '11', '12.9'],
       ['Peptic ulcer', '6', '7.1'],
       ['Normal', '10', '11.8'],
       ['Total', '85', '100.0']],
      'Each patient was assigned a single predominant endoscopic diagnosis. Any abnormality was '
      'present in 75 of 85 patients (88.2%).')

figure('FIGURE 1: Endoscopic appearance of esophagitis',
       'Esophagogastroduodenoscopy in a study participant showing mucosal breaks in the distal '
       'esophagus consistent with reflux esophagitis. Esophagitis was the preoperative finding in '
       '11 of 85 patients (12.9%).',
       '[PLACE FIGURE 1 IMAGE HERE - supplied separately as PNG]')

figure('FIGURE 2: Endoscopic appearance of a peptic ulcer',
       'Esophagogastroduodenoscopy in a study participant showing a discrete mucosal ulcer. '
       'Peptic ulcer disease was the preoperative finding in six of 85 patients (7.1%).',
       '[PLACE FIGURE 2 IMAGE HERE - supplied separately as PNG]')

figure('FIGURE 3: Endoscopic appearance of gastritis',
       'Esophagogastroduodenoscopy in a study participant showing erythematous, congested gastric '
       'mucosa consistent with gastritis. Gastritis was the commonest preoperative finding, '
       'present in 46 of 85 patients (54.1%).',
       '[PLACE FIGURE 3 IMAGE HERE - supplied separately as PNG]')

query('Cureus does not accept stretched, skewed or blurry images, and requires that any '
      'identifying patient information be removed from endoscopic images before upload. Upload '
      'each PNG through the Insert Figure dialog at the position marked above, entering the '
      'title and legend exactly as given.')

h2('Predictors of preoperative endoscopic abnormality')
p('Neither age, sex, nor duration of symptoms distinguished patients with an abnormal '
  'preoperative endoscopy from those with a normal examination. Mean age was 45.93 years '
  '(standard deviation 18.24) in the abnormal group and 46.20 years (standard deviation 14.14) in '
  'the normal group (p = 0.965). An abnormal endoscopy was found in 52 of 60 female patients '
  '(86.7%) and 23 of 25 male patients (92.0%) (p = 0.716). Mean symptom duration was 9.17 weeks '
  '(standard deviation 5.16) in the abnormal group and 9.50 weeks (standard deviation 6.21) in '
  'the normal group (p = 0.855). These comparisons are shown in Table 4.')

table('TABLE 4: Age, sex, and symptom duration according to preoperative endoscopic status',
      [['Variable', 'Abnormal endoscopy (n = 75)', 'Normal endoscopy (n = 10)', 'p value'],
       ['Mean age, years (SD)', '45.93 (18.24)', '46.20 (14.14)', '0.965'],
       ['Mean symptom duration, weeks (SD)', '9.17 (5.16)', '9.50 (6.21)', '0.855'],
       ['Female (n = 60), n (%)', '52 (86.7)', '8 (13.3)', '0.716'],
       ['Male (n = 25), n (%)', '23 (92.0)', '2 (8.0)', '']],
      'Continuous variables compared using independent-samples t-tests; sex compared using '
      "Fisher's exact test. Percentages in the sex rows are row percentages of the 60 female and "
      '25 male patients respectively. SD: standard deviation.',
      bold_last_row=False)

h2('Operative details')
p('Cholecystectomy was completed laparoscopically in 78 patients (91.8%), with conversion to open '
  'cholecystectomy in seven (8.2%).')
query('Insert here the exact number and nature of intraoperative complications, and the mean '
      '(standard deviation) postoperative hospital stay in days, taken from the master chart. The '
      'original wording ("complications were infrequent and minor", "approximately three days") '
      'has been removed because unquantified claims of this kind attract reviewer criticism and '
      'cannot be verified from the tables supplied.')

h2('Symptomatic outcome after operation')
p('One month after cholecystectomy, 64 patients (75.3%) were free of all six index symptoms and '
  '21 (24.7%) reported at least one persistent symptom.')
p('Symptom resolution was significantly more frequent among patients whose preoperative endoscopy '
  'had been abnormal and had therefore prompted medical treatment alongside surgery. Sixty of 75 '
  'patients (80.0%) in the abnormal group became symptom-free, against four of 10 (40.0%) in the '
  'normal group. The odds of symptom resolution were six times higher after an abnormal '
  'preoperative endoscopy than after a normal one (odds ratio 6.00, 95% confidence interval '
  "1.50-23.99; Fisher's exact test, p = 0.013). These data are shown in Table 5.")

table('TABLE 5: Preoperative endoscopic status and postoperative symptom outcome',
      [['Preoperative endoscopy', 'Symptom-free, n (%)', 'Persistent symptoms, n (%)', 'Total'],
       ['Abnormal', '60 (80.0)', '15 (20.0)', '75'],
       ['Normal', '4 (40.0)', '6 (60.0)', '10'],
       ['Total', '64 (75.3)', '21 (24.7)', '85']],
      "Fisher's exact test (two-tailed), p = 0.013. Odds ratio 6.00, 95% confidence interval "
      '1.50-23.99. Patients with abnormal preoperative findings received medical therapy for the '
      'lesion in addition to cholecystectomy. Row percentages are of the row total.')

h2('Endoscopic outcome after operation')
p('Repeat endoscopy at one month was compared with the preoperative examination. Among the 75 '
  'patients whose preoperative endoscopy had been abnormal, the appearances improved in 60 '
  '(80.0%), were unchanged in 11 (14.7%), and had worsened in four (5.3%). All 10 patients whose '
  'preoperative examination had been normal had a normal repeat examination and were therefore '
  'classified as unchanged. Endoscopic outcome was significantly associated with preoperative '
  'endoscopic status (Fisher-Freeman-Halton exact test, p < 0.001). Considered across the whole '
  'cohort, 15 patients (17.6%) had an abnormal endoscopy at one month and 70 (82.4%) a normal '
  'one. These data are shown in Table 6.')

table('TABLE 6: Preoperative endoscopic status and postoperative endoscopic outcome',
      [['Preoperative endoscopy', 'Improved, n (%)', 'Unchanged, n (%)', 'Worsened, n (%)',
        'Total'],
       ['Abnormal', '60 (80.0)', '11 (14.7)', '4 (5.3)', '75'],
       ['Normal', '0 (0.0)', '10 (100.0)', '0 (0.0)', '10'],
       ['Total', '60 (70.6)', '21 (24.7)', '4 (4.7)', '85']],
      'Fisher-Freeman-Halton exact test (two-tailed), p < 0.001. Endoscopic outcome at one month '
      'was classified relative to the preoperative examination. A preoperative examination that '
      'was normal and remained normal is classified as unchanged, because improvement relative to '
      'a normal baseline is not definable. Row percentages are of the row total.')

h2('Residual endoscopic pathology and persistent symptoms')
p('Persistence of symptoms at one month was closely related to the presence of residual mucosal '
  'disease at the same assessment. Of the 15 patients whose postoperative endoscopy remained '
  'abnormal, 12 (80.0%) were still symptomatic, whereas nine of the 70 patients (12.9%) whose '
  'postoperative endoscopy was normal reported persistent symptoms (odds ratio 27.11, 95% '
  "confidence interval 6.39-115.10; Fisher's exact test, p < 0.001). These data are shown in "
  'Table 7.')

table('TABLE 7: Postoperative endoscopic status and persistence of symptoms',
      [['Postoperative endoscopy', 'Persistent symptoms, n (%)', 'Symptom-free, n (%)', 'Total'],
       ['Abnormal', '12 (80.0)', '3 (20.0)', '15'],
       ['Normal', '9 (12.9)', '61 (87.1)', '70'],
       ['Total', '21 (24.7)', '64 (75.3)', '85']],
      "Fisher's exact test (two-tailed), p < 0.001. Odds ratio 27.11, 95% confidence interval "
      '6.39-115.10. Postoperative endoscopy was classified as normal or abnormal irrespective of '
      'the preoperative finding. Row percentages are of the row total.')

# ==================== DISCUSSION ====================
h1('DISCUSSION')
p('This study set out to establish how often patients presenting for cholecystectomy with upper '
  'gastrointestinal symptoms carry coexisting mucosal disease, and whether that disease matters '
  'to their outcome. Three findings emerged. Coexisting pathology was present in 88.2% of '
  'patients and could not be predicted from age, sex, or symptom duration. It was largely '
  'reversible, improving in four of every five affected patients after treatment and operation. '
  'And the patients least likely to become symptom-free were not those with the most pathology, '
  'but those with none.')
p('The demographic profile of the cohort was unremarkable and consistent with the established '
  'epidemiology of gallstone disease. A mean age in the fifth decade and a female predominance of '
  'approximately 2.4:1 correspond closely to figures reported from comparable Indian series by '
  'Gupta et al. and Brahmbhatt et al. [14,15]. This concordance suggests the cohort is '
  'representative of the population presenting for cholecystectomy in Indian government '
  'hospitals, and supports cautious generalization to comparable centers.')
p('The prevalence of abnormal preoperative endoscopy in this series, 88.2%, sits at the upper end '
  'of the published range. Reported yields span roughly 30% to 90%, and the spread is '
  'instructive. Series from European and South African centers have generally reported lower '
  'figures: Rashid et al. found abnormalities in 35% and Morrison and Mokoena in 35.4% [10,11]. '
  'Series from South Asian centers have reported considerably higher figures: 75.5% in the '
  '400-patient series of Kunnuru et al., 76.25% in the series of Khedkar et al., and 88.6% in the '
  'recent report of Maheshwari and Gond [12,16,17]. The present figure aligns with the South '
  'Asian pattern. Two explanations are likely, and they are not mutually exclusive. The first is '
  'a genuine population difference in the background burden of gastroduodenal mucosal disease, '
  'which in India reflects a high prevalence of *Helicobacter pylori* infection as well as '
  'dietary factors and unregulated non-steroidal anti-inflammatory drug use [13]. The second is '
  'diagnostic threshold: mild, non-erosive endoscopic gastritis is a subjective finding, and '
  'centers differ in whether they record it as pathology. Any comparison of raw yields across '
  'studies should be read with that caveat, and the clinically meaningful question is not how '
  'many patients have some abnormality but how many have an abnormality that alters management.')
p('Gastritis was the commonest lesion, in 54.1% of patients, consistent with essentially every '
  'published series on this question, including those of Gupta et al., Nasaruddin et al., Kolla '
  'et al., and Karmacharya et al. [14,18,19,20]. Peptic ulcer disease, found in six patients, and '
  'esophagitis, found in 11, carry importance disproportionate to their frequency, because each '
  'is an independent and treatable cause of the exact symptom complex that prompted referral, and '
  'each would have gone undetected on ultrasonography alone.')
p('That age, sex, and symptom duration all failed to predict an abnormal endoscopy is a '
  'practically significant negative result. It means the patients who harbor coexisting pathology '
  'cannot be picked out on clinical or demographic grounds at the bedside, an observation also '
  'made by Morrison and Mokoena and by Nasaruddin et al. [11,18]. If such patients are to be '
  'identified before operation, they must be identified endoscopically.')
p('The central finding of this study is the significant association between preoperative '
  'endoscopic status and symptomatic outcome, with an odds ratio of 6.00 favoring symptom '
  'resolution in patients whose endoscopy had been abnormal. The direction of this association '
  'warrants careful interpretation, because it runs opposite to that reported by some earlier '
  'authors. Nasaruddin et al. found that postoperative pain was significantly more common in '
  'patients with abnormal preoperative endoscopy [18]. Here, patients with an abnormal '
  'preoperative endoscopy were the more likely to become symptom-free, at 80.0% against 40.0%. '
  'The apparent contradiction resolves once the intervention is taken into account. In this '
  'study, mucosal lesions identified before operation were treated medically alongside surgery. A '
  'patient with treated gastritis therefore had both potential sources of symptoms addressed, '
  'whereas a patient with a normal endoscopy had only the gallbladder addressed, and if the '
  'symptoms were not in fact biliary, nothing had been done for them.')
p('The practical implication is that a normal preoperative endoscopy in a symptomatic patient '
  'with gallstones is not reassuring. It identifies a patient in whom the symptoms have no '
  'demonstrable mucosal cause, who may have functional dyspepsia or another non-biliary disorder '
  'that cholecystectomy will not relieve, and who therefore warrants a frank preoperative '
  'conversation about the likelihood of symptomatic benefit. In this cohort, only 40% of such '
  'patients were symptom-free at one month. This reading is consistent with Rashid et al., who '
  'reported that symptom persistence fell sharply when lesions were detected and treated before '
  'operation [10], and with the broader literature on post-cholecystectomy syndrome, in which the '
  'most treatable causes are those specifically identified and named rather than attributed to '
  'the syndrome label [9].')
p('The same principle is visible in the postoperative data, and in a still stronger form. '
  'Symptoms persisting at one month tracked residual mucosal disease closely: 12 of the 15 '
  'patients whose repeat endoscopy remained abnormal were still symptomatic, against only nine of '
  'the 70 whose repeat endoscopy was normal, an odds ratio of 27.11 whose confidence interval '
  'excludes unity by a wide margin. Symptoms after cholecystectomy, in other words, followed the '
  'state of the upper gastrointestinal mucosa rather than the biliary operation. The clinical '
  'message is direct. When a patient complains of upper gastrointestinal symptoms after '
  'cholecystectomy, the probability that mucosal disease is responsible is high, and the '
  'appropriate response is endoscopic assessment rather than attribution to post-cholecystectomy '
  'syndrome as a diagnosis of exclusion. Girometti et al. made a parallel argument from the '
  'imaging side, showing that a structured search identifies specific biliary abnormalities in '
  'patients carrying that label [21], and Ruiz-Campos et al. demonstrated that a further specific '
  'and treatable cause, bile acid malabsorption, accounts for a large share of chronic diarrhea '
  'after cholecystectomy [22]. The common theme is that the syndrome label conceals treatable '
  'diagnoses.')
p('Some caution is nonetheless warranted. The comparison rests on only 10 patients with a normal '
  'preoperative endoscopy, and the confidence interval around the odds ratio of 6.00 is '
  'correspondingly wide, running from 1.50 to 23.99. The direction and significance of the effect '
  'are established, but its magnitude is estimated imprecisely and should be confirmed in a '
  'larger series before being used for individual prognostication.')
p('The parallel improvement in endoscopic findings, seen in 60 of the 75 patients with abnormal '
  'preoperative examinations (80.0%), deserves comment. It should not be read as evidence that '
  'cholecystectomy heals gastritis. These patients received medical therapy directed at the '
  'lesions found, and the expected response of gastritis, duodenitis, and esophagitis to acid '
  'suppression over a month is precisely the improvement observed. The correct inference is '
  'narrower but still useful: mucosal disease coexisting with gallstones is largely reversible '
  'over a short interval when it is looked for and treated. The four patients whose findings '
  'worsened are a reminder that the postoperative period is not uniformly benign, and are '
  'consistent with reports of altered bile flow and bile acid handling after gallbladder removal '
  '[22].')
p('Whether these findings justify endoscopy for every patient, as against a low threshold for '
  'endoscopy in selected patients, is a question this study cannot fully answer, because every '
  'patient in the cohort had symptoms and every patient was scoped. There was no unscoped '
  'comparison group. In a population with an 88.2% yield of abnormal findings, a policy of '
  'routine endoscopy for symptomatic patients is defensible on diagnostic grounds; whether it is '
  'defensible on economic grounds depends on local endoscopy capacity and cost, which this study '
  'did not measure. Kunnuru et al. reached the position that routine scopy is not necessary for '
  'every patient but is particularly valuable in those with atypical presentations [12]. The same '
  'logic has been applied in the analogous setting of suspected choledocholithiasis, where Maple '
  'et al. showed that risk-stratified rather than universal endoscopic evaluation optimizes '
  'diagnostic yield [23]. The present data are compatible with that position, given that all '
  'patients here were symptomatic by definition.')

h2('Limitations')
p('This study has several limitations. All patients underwent endoscopy and all received '
  'treatment for lesions found, so the effect of endoscopy cannot be separated from the effect of '
  'the medical therapy it prompted, and no inference about the counterfactual of not scoping is '
  'possible. The normal-endoscopy comparison group comprised only 10 patients, which limits the '
  'precision of the principal estimate and means that comparisons involving that group rest on '
  'small cell counts. Follow-up extended to one month only, adequate for mucosal healing but too '
  'short to characterize late symptom recurrence or to capture the delayed presentations that '
  'constitute much of post-cholecystectomy syndrome. Symptoms were recorded as present or absent '
  'rather than graded on a validated dyspepsia severity instrument, which limits sensitivity to '
  'partial improvement and precludes comparison with studies using scored outcomes. '
  '*Helicobacter pylori* status was not determined, a material omission given that it is the '
  'leading candidate explanation for the high gastritis prevalence observed and would have '
  'permitted a directly actionable eradication analysis. Endoscopic findings were assigned to a '
  'single predominant category per patient, so coexisting lesions in the same patient are not '
  'represented. Patients with complicated gallstone disease were excluded, so these results do '
  'not extend to that group. Finally, the study was conducted at a single center using '
  'convenience sampling, and the 88.2% prevalence should not be transported to populations with a '
  'different background burden of gastroduodenal disease.')
p('Future work should compare scoped and unscoped cohorts prospectively, recruit a larger '
  'normal-endoscopy group, incorporate *Helicobacter pylori* testing and eradication, use '
  'validated symptom scores, and extend follow-up to at least six to 12 months. A formal '
  'cost-effectiveness analysis in a resource-constrained setting would be of particular value, '
  'since the case for routine preoperative endoscopy in such settings ultimately turns on whether '
  'the cost of scoping all symptomatic patients is less than the cost of investigating and '
  're-treating those who remain symptomatic afterwards.')

# ==================== CONCLUSIONS ====================
h1('CONCLUSIONS')
p('Coexisting upper gastrointestinal mucosal disease was present in almost 90% of patients '
  'presenting with symptomatic cholelithiasis in this cohort, with gastritis accounting for the '
  'majority and peptic ulcer and esophagitis present in a clinically important minority. This '
  'pathology could not be predicted from age, sex, or duration of symptoms, and would not have '
  'been detected by the ultrasonography that established the diagnosis of gallstones. When '
  'identified before operation and treated medically alongside cholecystectomy, it proved largely '
  'reversible, with endoscopic findings improving in four of every five affected patients within '
  'a month.')
p('The value of preoperative endoscopy in this setting therefore lies less in changing the '
  'decision to operate than in disclosing a second, treatable source of symptoms that would '
  'otherwise be attributed to the gallbladder and left untreated. The corollary is equally '
  'important for counseling: patients whose preoperative endoscopy is normal were significantly '
  'less likely to become symptom-free after cholecystectomy, and a normal examination should '
  'prompt a careful discussion of expected benefit rather than reassurance. Symptoms that '
  'persisted a month after operation were, in turn, overwhelmingly confined to the patients in '
  'whom mucosal disease had persisted rather than resolved, which argues for endoscopic '
  'reassessment of such patients rather than attribution to post-cholecystectomy syndrome without '
  'further evaluation. A low threshold for esophagogastroduodenoscopy before cholecystectomy is '
  'accordingly justified in patients presenting with dyspeptic or atypical upper gastrointestinal '
  'symptoms, particularly where the background prevalence of gastroduodenal disease is high.')

# ==================== ADDITIONAL INFORMATION ====================
h1('ADDITIONAL INFORMATION')

h2('Author contributions')
p('[REDACTED FOR BLINDED REVIEW. In the unblinded submission copy, list concept and design, '
  'acquisition/analysis/interpretation of data, drafting of the manuscript, critical review, and '
  'supervision against the named authors, as required by Cureus.]')

h2('Disclosures')
p('Human subjects: Consent was obtained or waived by all participants in this study. The '
  'institutional Scientific Review Committee and the Human Research Ethics Committee of the '
  'participating institution [name redacted for blinded review] issued approval [reference number '
  'redacted for blinded review].')
p('Animal subjects: All authors have confirmed that this study did not involve animal subjects or '
  'tissue.')
p('Conflicts of interest: In compliance with the ICMJE uniform disclosure form, all authors '
  'declare the following:')
p('Payment/services info: All authors have declared that no financial support was received from '
  'any organization for the submitted work.')
p('Financial relationships: All authors have declared that they have no financial relationships '
  'at present or within the previous three years with any organizations that might have an '
  'interest in the submitted work.')
p('Other relationships: All authors have declared that there are no other relationships or '
  'activities that could appear to have influenced the submitted work.')
query('Insert the Human Research Ethics Committee approval reference number and date in the '
      'unblinded submission copy. Cureus will not publish without it.')

# ==================== REFERENCES ====================
h1('REFERENCES')
refs = [
 'Wang X, Yu W, Jiang G, et al.: Global epidemiology of gallstones in the 21st century: a '
 'systematic review and meta-analysis. Clin Gastroenterol Hepatol. 2024, 22:1586-95. '
 '10.1016/j.cgh.2024.01.051',
 'Unalp-Arida A, Ruhl CE: Burden of gallstone disease in the United States population: '
 'prepandemic rates and trends. World J Gastrointest Surg. 2024, 16:1130-48. '
 '10.4240/wjgs.v16.i4.1130',
 'Sun H, Warren J, Yip J, Ji Y, Hao S, Han W, Ding Y: Factors influencing gallstone formation: a '
 'review of the literature. Biomolecules. 2022, 12:550. 10.3390/biom12040550',
 'Di Ciaula A, Wang DQ, Portincasa P: An update on the pathogenesis of cholesterol gallstone '
 'disease. Curr Opin Gastroenterol. 2018, 34:71-80. 10.1097/MOG.0000000000000423',
 'Yuk JS, Park JY: Menopausal hormone therapy increases the risk of gallstones: Health Insurance '
 'Database in South Korea (HISK)-based cohort study. PLoS One. 2023, 18:e0294356. '
 '10.1371/journal.pone.0294356',
 'Fujita N, Yasuda I, Endo I, et al.: Evidence-based clinical practice guidelines for '
 'cholelithiasis 2021. J Gastroenterol. 2023, 58:801-33. 10.1007/s00535-023-02014-6',
 'European Association for the Study of the Liver (EASL): EASL Clinical Practice Guidelines on '
 'the prevention, diagnosis and treatment of gallstones. J Hepatol. 2016, 65:146-81. '
 '10.1016/j.jhep.2016.03.005',
 'Shabanzadeh DM: The symptomatic outcomes of cholecystectomy for gallstones. J Clin Med. 2023, '
 '12:1897. 10.3390/jcm12051897',
 'Nam C, Lee JS, Kim JS, Lee TY, Yoon YC: Clinical perspectives on post-cholecystectomy syndrome: '
 'a narrative review. Ann Med. 2025, 57:2496408. 10.1080/07853890.2025.2496408',
 'Rashid F, Rashid N, Waraich N, Ahmed J, Iftikhar SY: Role of routine oesophago-gastroduodenoscopy '
 'before cholecystectomy. Int J Surg. 2010, 8:236-8. 10.1016/j.ijsu.2010.01.008',
 'Morrison S, Mokoena T: Routine upper gastro-intestinal tract endoscopy before elective '
 'cholecystectomy for symptomatic gallstones-justified. Sci Rep. 2024, 14:14042. '
 '10.1038/s41598-024-64019-2',
 'Kunnuru SKR, Kanmaniyan B, Thiyagarajan M, Singh BK, Navrathan N: A study on efficacy of UGI '
 'scopy in cholelithiasis patients before laparoscopic cholecystectomy. Minim Invasive Surg. '
 '2021, 2021:8849032. 10.1155/2021/8849032',
 'Romshoo GJ, Malik GM, Bhat MY, Rather AR, Basu JA, Qureshi KA: *Helicobacter pylori* associated '
 'antral gastritis in peptic ulcer disease patients and normal healthy population of Kashmir, '
 'India. Diagn Ther Endosc. 1998, 4:135-9. 10.1155/DTE.4.135',
 'Gupta P, Gupta V, Singh SP, et al.: Role of routine upper gastro intestinal endoscopy in '
 'patients of cholelithiasis presenting with dyspepsia in rural set-up. Int Surg J. 2016, '
 '3:509-15. 10.18203/2349-2902.isj20160951',
 'Brahmbhatt H, Pundeer S, Bhatia P, Goel K, Garg U: Diagnostic value of upper gastrointestinal '
 'endoscopy prior to cholecystectomy in a tertiary care institute. J Evol Med Dent Sci. 2020, '
 '9:817-21. 10.14260/jemds/2020/177',
 'Khedkar I, Prasad D, Datta A: Diagnostic value of upper gastrointestinal endoscopy prior to '
 'elective laparoscopic cholecystectomy for symptomatic cholelithiasis. Int Surg J. 2018, 5:105-9. '
 '10.18203/2349-2902.isj20175549',
 'Maheshwari N, Gond P: Study of upper GI endoscopy findings and plan of treatment in patients '
 'with cholelithiasis. Int J Sci Res. 2025, 14:997-1001. 10.21275/sr25215193944',
 'Nasaruddin A, Jain D, Patil K, Phalgune D: Diagnostic and prognostic role of upper '
 'gastrointestinal endoscopy in cholelithiasis patients with upper gastrointestinal symptoms. '
 'Saudi Surg J. 2020, 8:185-91. 10.4103/ssj.ssj_84_21',
 'Kolla V, Charles N, Datey S, Mahor D, Gupta A, Malhotra S: Upper gastrointestinal endoscopy '
 'prior to laparoscopic cholecystectomy: a clinical study at a tertiary care centre in central '
 'India. Int Surg J. 2016, 3:637-42. 10.18203/2349-2902.isj20161136',
 'Karmacharya A, Malla BR, Joshi HN, Gurung RB, Rajbhandari M: The predictive value of '
 'pre-operative symptoms including upper gastrointestinal endoscopy before laparoscopic '
 'cholecystectomy for elective symptomatic cholecystolithiasis. Kathmandu Univ Med J (KUMJ). '
 '2013, 11:300-4. 10.3126/kumj.v11i4.12526',
 'Girometti R, Brondani G, Cereser L, Como G, Del Pin M, Bazzocchi M, Zuiani C: '
 'Post-cholecystectomy syndrome: spectrum of biliary findings at magnetic resonance '
 'cholangiopancreatography. Br J Radiol. 2010, 83:351-61. 10.1259/bjr/99865290',
 'Ruiz-Campos L, Gisbert JP, Ysamat M, Arau B, Loras C, Esteve M, Fernandez-Banares F: Systematic '
 'review with meta-analysis: the prevalence of bile acid malabsorption and response to '
 'colestyramine in patients with chronic watery diarrhoea and previous cholecystectomy. Aliment '
 'Pharmacol Ther. 2019, 49:242-50. 10.1111/apt.15099',
 'Maple JT, Ben-Menachem T, Anderson MA, et al.: The role of endoscopy in the evaluation of '
 'suspected choledocholithiasis. Gastrointest Endosc. 2010, 71:1-9. 10.1016/j.gie.2009.09.041',
]
for i, r in enumerate(refs, 1):
    par = doc.add_paragraph()
    runs(par, f'{i}. {r}', size=11)
    par.paragraph_format.space_after = Pt(5)
    par.paragraph_format.line_spacing = 1.15

doc.save('blinded_manuscript.docx')
print('saved; references:', len(refs))

# --- post-process: fix python-docx's schema-invalid <w:zoom> in settings.xml ---
import zipfile, shutil, os
src = 'blinded_manuscript.docx'
tmp = 'blinded_manuscript.fixed.docx'
zin = zipfile.ZipFile(src)
zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == 'word/settings.xml':
        data = data.replace(b'<w:zoom w:percent="100"/>', b'')
        data = data.replace(b'<w:zoom/>', b'')
        data = re.sub(rb'<w:zoom(?![a-zA-Z])[^>]*/>', b'', data)
        data = re.sub(rb'<w:zoom(?![a-zA-Z])[^>]*>.*?</w:zoom>', b'', data, flags=re.S)
    zout.writestr(item, data)
zout.close(); zin.close()
shutil.move(tmp, src)
print('settings.xml patched')
